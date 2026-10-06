import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from pylibs.core.instrument_response import (
    InstrumentResponseLibrary,
    ResponseCurve,
    extract_window_tag,
    parse_lsf_widths,
    parse_response_paths,
    serialize_lsf_widths,
    serialize_response_paths,
)
from pyLIBS import (
    AppOptions,
    MainWindow,
    OptionsWindow,
    Spectrum,
    SpectrumSourceWindow,
    merge_spectra_by_wavelength,
    merged_spectrum_display_name,
    opc_calibration_id,
    opc_rows_calibration_id,
    read_opc_metadata,
    reset_measurement_session,
    select_spectrum_for_wavelengths,
    spectrum_plot_segments,
    spectrum_covers_wavelengths,
    spectrum_wavelength_intervals,
    wavelength_is_acquired,
)


class InstrumentResponseTests(unittest.TestCase):
    def test_opc_identity_distinguishes_regenerated_calibrations(self):
        rows_a = [(313.411, 4.0, "Ni", "Ni I"), (340.365, 5.0, "Cd", "Cd I")]
        rows_b = [(313.411, 4.1, "Ni", "Ni I"), (340.365, 5.0, "Cd", "Cd I")]

        self.assertNotEqual(
            opc_rows_calibration_id(rows_a),
            opc_rows_calibration_id(rows_b),
        )

    def test_legacy_opc_gets_stable_file_identity_and_reads_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "standard.opc"
            path.write_text(
                "# pyLIBS One Point Calibration\n"
                "# source spectrum: standard.txt\n"
                "# nominal_number_percent: {\"Ni\": 60.0}\n"
                "313.411 4.0 Ni Ni I\n"
                "340.365 5.0 Cd Cd I\n",
                encoding="utf-8",
            )

            first = opc_calibration_id(str(path))
            second = opc_calibration_id(str(path))
            metadata = read_opc_metadata(str(path))

            self.assertEqual(first, second)
            self.assertEqual(len(first), 12)
            self.assertEqual(metadata["source_spectrum"], "standard.txt")
            self.assertEqual(metadata["nominal_number_percent"], {"Ni": 60.0})

    def test_new_measurement_clears_sample_state_but_can_label_series_ne(self):
        app = SimpleNamespace(
            options=SimpleNamespace(remember_ne_for_series=True),
            session_ne=1.23e17,
            session_ne_source="H-alpha 656.txt",
            last_known_ne=None,
            last_known_ne_source=None,
            session_temperature=0.9,
            session_temperature_source="Saha-Boltzmann",
            post_irsac_temperature=0.8,
            post_irsac_temperature_source="IRSAC",
            last_halpha_result={"ne": 1.23e17},
            last_saha_boltzmann={"kt": 0.9},
            last_cflibs_rows=[{"element": "Ni"}],
            last_cflibs_quality={"status": "PASS"},
            last_cflibs_robustness={"status": "PASS"},
            opc_nominal_percent={"Ni": 100.0},
            opc_nominal_inputs={"mass_percent": {"Ni": 100.0}},
            session_voigt_fit_range_keys={"Ni:1:313.411000"},
        )

        message = reset_measurement_session(app)

        self.assertIsNone(app.session_temperature)
        self.assertEqual(app.last_cflibs_rows, [])
        self.assertEqual(app.opc_nominal_percent, {})
        self.assertEqual(app.session_voigt_fit_range_keys, set())
        self.assertAlmostEqual(app.session_ne, 1.23e17)
        self.assertIn("series carry-over", app.session_ne_source)
        self.assertIn("Ne retained for series", message)

    def test_ascii_rounding_at_endpoint_uses_endpoint_value(self):
        curve = ResponseCurve(
            [313.41101, 479.99121],
            [434357.7, 482332.42],
            "standard.opc",
        )

        self.assertAlmostEqual(curve.value_at(479.99121094), 482332.42)
        self.assertIsNone(curve.value_at(479.992))

    def test_selecting_physical_response_preserves_apply_after_choice(self):
        class Var:
            def __init__(self, value=None):
                self.value = value

            def get(self):
                return self.value

            def set(self, value):
                self.value = value

        path = r"C:\response\instrument.xlsx"
        window = SimpleNamespace(
            master_app=SimpleNamespace(options=SimpleNamespace()),
            vars={
                "response_file": Var(""),
                "response_files": Var(""),
                "apply_response": Var(False),
                "global_correction_mode": Var("after"),
            },
            _response_selection_pending=False,
            _set_response_summary=lambda paths: None,
        )

        with patch("pyLIBS.filedialog.askopenfilenames", return_value=(path,)), patch(
            "pyLIBS.remember_working_dir"
        ):
            OptionsWindow.browse_response(window)

        self.assertEqual(window.vars["global_correction_mode"].get(), "after")
        self.assertTrue(window.vars["apply_response"].get())
        self.assertTrue(window._response_selection_pending)

    def test_options_reloads_a_reselected_response_through_the_main_loader(self):
        path = r"C:\response\instrument.xlsx"
        options = SimpleNamespace(
            response_file=path,
            response_files=serialize_response_paths([path]),
        )
        calls = []
        app = SimpleNamespace(
            options=options,
            configured_response_paths=lambda: parse_response_paths(options.response_files),
            load_response_files=lambda paths: calls.append(list(paths)) or True,
        )
        window = SimpleNamespace(
            master_app=app,
            _response_selection_pending=True,
        )

        loaded = OptionsWindow.reload_response_if_changed(window, [path])

        self.assertTrue(loaded)
        self.assertEqual(calls, [[path]])
        self.assertFalse(window._response_selection_pending)

    def test_options_keeps_selection_pending_when_response_loading_fails(self):
        previous = r"C:\response\previous.xlsx"
        selected = r"C:\response\invalid.xlsx"
        options = SimpleNamespace(
            response_file=selected,
            response_files=serialize_response_paths([selected]),
        )
        app = SimpleNamespace(
            options=options,
            configured_response_paths=lambda: parse_response_paths(options.response_files),
            load_response_files=lambda paths: False,
        )
        window = SimpleNamespace(
            master_app=app,
            _response_selection_pending=True,
        )

        loaded = OptionsWindow.reload_response_if_changed(window, [previous])

        self.assertFalse(loaded)
        self.assertTrue(window._response_selection_pending)

    def test_response_workbook_expands_windows_and_uses_scaled_tungsten_column(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "instrument_response.xlsx"
            workbook = openpyxl.Workbook()
            deuterium = workbook.active
            deuterium.title = "fs365d"
            deuterium.append([344.0, 10.0])
            deuterium.append([345.0, 20.0])
            overlap = workbook.create_sheet("fs365w")
            overlap.append([344.0, 0.1, 100.0])
            overlap.append([345.0, 0.2, 200.0])
            tungsten = workbook.create_sheet("fs400w")
            tungsten.append([390.0, 0.3, 300.0])
            tungsten.append([410.0, 0.4, 400.0])
            workbook.save(path)

            library = InstrumentResponseLibrary.from_files([str(path)])

        self.assertEqual(len(library.curves), 2)
        by_tag = {extract_window_tag(curve.name): curve for curve in library.curves}
        self.assertEqual(by_tag["365"].y, [10.0, 20.0])
        self.assertEqual(by_tag["400"].y, [300.0, 400.0])

    def test_paths_round_trip(self):
        paths = [r"C:\response\300 nm.txt", r"C:\response\400 nm.txt"]
        self.assertEqual(parse_response_paths(serialize_response_paths(paths)), paths)

    def test_lsf_widths_round_trip_and_per_window_lookup(self):
        widths = {"fs305.txt": 0.12, "fs475.txt": 0.25}
        self.assertEqual(parse_lsf_widths(serialize_lsf_widths(widths)), widths)
        curve = ResponseCurve([300.0, 310.0], [1.0, 1.0], "fs305.txt", 0.12)
        library = InstrumentResponseLibrary([curve])
        self.assertAlmostEqual(library.lsf_fwhm_at(305.0, 0.5), 0.12)
        self.assertAlmostEqual(library.lsf_fwhm_at(500.0, 0.5), 0.5)

    def test_match_uses_window_range(self):
        low = ResponseCurve([300, 305, 310], [1, 1, 1], "response 300 nm.txt")
        high = ResponseCurve([400, 405, 410], [2, 2, 2], "response 400 nm.txt")
        library = InstrumentResponseLibrary([low, high])
        self.assertIs(library.best_for("spectrum 400 nm.txt", [401, 405, 409]), high)
        self.assertIs(library.best_for("unknown.txt", [301, 305, 309]), low)

    def test_outside_response_is_not_clamped(self):
        curve = ResponseCurve([300, 305, 310], [1, 2, 3], "response.txt")
        self.assertIsNone(curve.value_at(299.9))
        self.assertIsNone(curve.value_at(310.1))

    def test_overlapping_windows_get_distinct_tagged_curves(self):
        first = ResponseCurve([300, 305, 310], [1, 1, 1], "response 300 nm.txt")
        second = ResponseCurve([300, 305, 310], [2, 2, 2], "response 310 nm.txt")
        library = InstrumentResponseLibrary([first, second])
        assignments = library.assign_windows([
            ("spectrum 300 nm.txt", [301, 305, 309]),
            ("spectrum 310 nm.txt", [301, 305, 309]),
        ])
        self.assertIs(assignments[0], first)
        self.assertIs(assignments[1], second)

    def test_correction_trims_unreliable_edges(self):
        x = [300.0 + index * 0.1 for index in range(101)]
        response = [0.0001 if index < 10 or index > 90 else 2.0 for index in range(101)]
        curve = ResponseCurve(x, response, "response 305 nm.txt")
        corrected_x, corrected_y, info = curve.correct(x, [20.0] * len(x))
        self.assertAlmostEqual(corrected_x[0], 301.0)
        self.assertAlmostEqual(corrected_x[-1], 309.0)
        self.assertTrue(all(abs(value - 10.0) < 1e-12 for value in corrected_y))
        self.assertIn("edge threshold", info)

    def test_joined_windows_are_corrected_before_merging(self):
        x1 = [300.0 + index * 0.1 for index in range(51)]
        x2 = [304.0 + index * 0.1 for index in range(51)]
        first = Spectrum(x1, [10.0] * len(x1), name="window 300 nm")
        second = Spectrum(x2, [30.0] * len(x2), name="window 309 nm")
        joined = merge_spectra_by_wavelength(first, second, "joined")
        library = InstrumentResponseLibrary([
            ResponseCurve(x1, [2.0] * len(x1), "response 300 nm"),
            ResponseCurve(x2, [3.0] * len(x2), "response 309 nm"),
        ])
        app = SimpleNamespace(response_library=library)
        corrected = MainWindow._response_corrected_copy(app, joined)
        self.assertTrue(corrected.response_corrected)
        self.assertEqual(len(corrected.response_sources), 2)
        self.assertEqual(len(corrected.source_windows), 2)
        self.assertAlmostEqual(min(corrected.y), 5.0)
        self.assertAlmostEqual(max(corrected.y), 10.0)

    def test_response_uses_raw_axis_after_wavelength_calibration(self):
        raw_x = [300.0 + index * 0.1 for index in range(11)]
        calibrated_x = [value - 0.4 for value in raw_x]
        source = SpectrumSourceWindow(
            tuple(calibrated_x),
            tuple([10.0] * len(raw_x)),
            "window 300 nm",
            raw_x=tuple(raw_x),
            calibration_intercept=-0.4,
            calibration_slope=1.0,
            calibration_model="offset",
        )
        spectrum = Spectrum(
            calibrated_x,
            [10.0] * len(raw_x),
            name="calibrated",
            source_windows=(source,),
        )
        library = InstrumentResponseLibrary([
            ResponseCurve(raw_x, [2.0] * len(raw_x), "response 300 nm")
        ])
        app = SimpleNamespace(response_library=library)

        corrected = MainWindow._response_corrected_copy(app, spectrum)

        self.assertAlmostEqual(corrected.x[0], 299.6)
        self.assertAlmostEqual(corrected.x[-1], 300.6)
        self.assertTrue(all(abs(value - 5.0) < 1e-12 for value in corrected.y))

    def test_only_halpha_ne_loader_can_bypass_missing_radiometric_response(self):
        with tempfile.TemporaryDirectory() as tmp:
            spectrum_path = Path(tmp) / "halpha_656.txt"
            spectrum_path.write_text(
                "\n".join(
                    f"{650.0 + index * 0.5:.3f}\t{10.0 + index:.3f}"
                    for index in range(25)
                ),
                encoding="ascii",
            )
            options = AppOptions(
                apply_response=True,
                apply_before=True,
                auto_wavelength_calibration=False,
                dark_mode="constant",
                dark_constant=2.0,
            )
            response = ResponseCurve(
                [300.0, 305.0, 310.0],
                [1.0, 1.0, 1.0],
                "response 305 nm",
            )
            app = SimpleNamespace(
                options=options,
                response_library=InstrumentResponseLibrary([response]),
                apply_auto_offset_if_enabled=lambda spectrum: spectrum,
                apply_wavelength_calibration_if_enabled=lambda spectrum: spectrum,
            )
            app._response_corrected_copy = (
                lambda spectrum: MainWindow._response_corrected_copy(app, spectrum)
            )

            with self.assertRaisesRegex(ValueError, "No response curve matches"):
                MainWindow.load_spectrum_with_corrections(app, str(spectrum_path))

            halpha = MainWindow.load_halpha_spectrum_for_ne(app, str(spectrum_path))

            self.assertFalse(halpha.response_corrected)
            self.assertTrue(halpha.halpha_width_only)
            self.assertTrue(halpha.halpha_response_bypassed)
            self.assertEqual(len(halpha.x), 25)
            self.assertAlmostEqual(halpha.y[0], 8.0)

    def test_overlapping_windows_use_only_a_short_crossfade(self):
        first = Spectrum(list(range(9)), [10] * 9, name="lower")
        second = Spectrum(list(range(2, 11)), [20] * 9, name="upper")
        joined = merge_spectra_by_wavelength(first, second, "joined")
        values = dict(zip(joined.x, joined.y))
        self.assertAlmostEqual(values[4], 10.0)
        self.assertAlmostEqual(values[5], 15.0)
        self.assertAlmostEqual(values[6], 20.0)

    def test_display_merge_does_not_blend_a_peak_far_from_the_splice(self):
        first = Spectrum(list(range(11)), [0, 0, 0, 0, 0, 0, 0, 100, 0, 0, 0], name="lower")
        second = Spectrum(list(range(2, 13)), [0, 0, 0, 0, 0, 200, 0, 0, 0, 0, 0], name="upper")
        joined = merge_spectra_by_wavelength(first, second, "joined")
        values = dict(zip(joined.x, joined.y))
        self.assertAlmostEqual(values[7], 200.0)

    def test_display_merge_keeps_a_line_from_the_more_central_window(self):
        lower_x = [round(305.0 + 0.1 * index, 1) for index in range(401)]
        upper_x = [round(330.0 + 0.1 * index, 1) for index in range(401)]
        lower = Spectrum(lower_x, [1.0] * len(lower_x), name="325.txt")
        upper = Spectrum(
            upper_x,
            [10.0 if wavelength == 340.7 else 1.0 for wavelength in upper_x],
            name="350.txt",
        )

        joined = merge_spectra_by_wavelength(lower, upper, "joined")
        values = dict(zip(joined.x, joined.y))

        self.assertAlmostEqual(values[340.7], 10.0)

    def test_merged_analysis_uses_one_physical_source_window(self):
        first = Spectrum(list(range(7)), [10, 10, 10, 10, 100, 10, 10], name="350.txt")
        second = Spectrum(list(range(4, 11)), [20, 20, 200, 20, 20, 20, 20], name="365.txt")
        joined = merge_spectra_by_wavelength(first, second, "joined")
        app = SimpleNamespace(
            spectra=[joined],
            independent_window_mode=False,
            active_window=None,
        )

        selected = select_spectrum_for_wavelengths(app, [4.5])

        self.assertIsNot(selected, joined)
        self.assertEqual(selected.name, "350.txt")
        self.assertEqual(selected.y, first.y)

    def test_merged_analysis_rejects_a_fit_spanning_two_source_windows(self):
        first = Spectrum([0.0, 1.0, 2.0], [1.0, 2.0, 1.0], name="first")
        second = Spectrum([4.0, 5.0, 6.0], [1.0, 3.0, 1.0], name="second")
        joined = merge_spectra_by_wavelength(first, second, "joined")
        app = SimpleNamespace(
            spectra=[joined],
            independent_window_mode=False,
            active_window=None,
        )

        self.assertIsNone(select_spectrum_for_wavelengths(app, [1.0, 5.0]))

    def test_independent_mode_selects_one_real_window_without_merging(self):
        first = Spectrum(
            [300.0, 301.0, 302.0],
            [10.0, 11.0, 12.0],
            name="305.txt",
        )
        second = Spectrum(
            [301.0, 302.0, 303.0],
            [20.0, 21.0, 22.0],
            name="325.txt",
        )
        app = SimpleNamespace(
            spectra=[first, second],
            independent_window_mode=True,
            active_window=None,
        )

        selected = select_spectrum_for_wavelengths(app, [302.8])

        self.assertIs(selected, second)
        self.assertEqual(selected.y, [20.0, 21.0, 22.0])
        self.assertTrue(wavelength_is_acquired(app, 302.8))

    def test_independent_mode_never_bridges_two_detector_windows(self):
        first = Spectrum([300.0, 301.0], [1.0, 2.0], name="305.txt")
        second = Spectrum([302.0, 303.0], [3.0, 4.0], name="325.txt")
        app = SimpleNamespace(
            spectra=[first, second],
            independent_window_mode=True,
            active_window=None,
        )

        self.assertFalse(spectrum_covers_wavelengths(first, [300.5, 302.5]))
        self.assertIsNone(
            select_spectrum_for_wavelengths(app, [300.5, 302.5])
        )

    def test_active_window_resolves_an_overlap_in_independent_mode(self):
        first = Spectrum([300.0, 301.0, 302.0], [1.0, 1.0, 1.0], name="305.txt")
        second = Spectrum([301.0, 302.0, 303.0], [2.0, 2.0, 2.0], name="325.txt")
        active = SimpleNamespace(
            selected_index=1,
            winfo_exists=lambda: True,
        )
        app = SimpleNamespace(
            spectra=[first, second],
            independent_window_mode=True,
            active_window=active,
        )

        self.assertIs(select_spectrum_for_wavelengths(app, [301.5]), second)

    def test_non_overlapping_windows_are_not_connected_for_display(self):
        first = Spectrum([300.0, 300.1, 300.2], [1.0, 2.0, 1.0], name="first")
        second = Spectrum([400.0, 400.1, 400.2], [3.0, 4.0, 3.0], name="second")
        joined = merge_spectra_by_wavelength(first, second)
        segments = spectrum_plot_segments(joined)
        self.assertEqual(len(segments), 2)
        self.assertEqual(segments[0][0], first.x)
        self.assertEqual(segments[1][0], second.x)
        self.assertEqual(merged_spectrum_display_name(joined), "Merged spectrum (2 windows)")

    def test_acquired_coverage_excludes_unmeasured_gap_and_high_wavelengths(self):
        first = Spectrum([290.0, 300.0, 330.0], [1.0, 2.0, 1.0], name="first")
        second = Spectrum([350.0, 380.0, 410.0], [3.0, 4.0, 3.0], name="second")
        joined = merge_spectra_by_wavelength(first, second)
        app = SimpleNamespace(spectra=[joined])

        self.assertEqual(
            spectrum_wavelength_intervals(joined),
            [(290.0, 330.0), (350.0, 410.0)],
        )
        self.assertTrue(wavelength_is_acquired(app, 405.78))
        self.assertFalse(wavelength_is_acquired(app, 340.0))
        self.assertFalse(wavelength_is_acquired(app, 465.11))
        self.assertFalse(wavelength_is_acquired(app, 479.99))

    def test_overlapping_windows_are_one_display_segment(self):
        first = Spectrum([300.0, 301.0, 302.0], [1.0, 2.0, 1.0], name="first")
        second = Spectrum([301.0, 302.0, 303.0], [3.0, 4.0, 3.0], name="second")
        joined = merge_spectra_by_wavelength(first, second)
        self.assertEqual(len(spectrum_plot_segments(joined)), 1)

    def test_trimmed_response_gap_breaks_even_if_raw_windows_overlap(self):
        raw_first = SpectrumSourceWindow((300.0, 302.0, 304.0), (1.0, 1.0, 1.0), "first")
        raw_second = SpectrumSourceWindow((303.0, 305.0, 307.0), (1.0, 1.0, 1.0), "second")
        corrected = Spectrum(
            [300.0, 300.1, 300.2, 306.8, 306.9, 307.0],
            [1.0, 2.0, 1.0, 3.0, 4.0, 3.0],
            name="corrected",
            source_windows=(raw_first, raw_second),
        )
        self.assertEqual(len(spectrum_plot_segments(corrected)), 2)


if __name__ == "__main__":
    unittest.main()
